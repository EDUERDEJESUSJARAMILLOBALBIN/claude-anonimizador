"""Reading file formats into in-memory tables. Nothing read here leaves the engine."""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")  # Excel in Colombia exports CSV as cp1252
STRUCTURED = {"json": "json", "jsonl": "jsonl", "ndjson": "jsonl", "yaml": "yaml", "yml": "yaml"}


@dataclass
class Table:
    name: str | None
    columns: list[str]
    rows: list[list[object]]
    hidden: bool = False

    def values(self, index: int) -> list[object]:
        return [row[index] if index < len(row) else None for row in self.rows]


@dataclass
class Reading:
    file_format: str
    tables: list[Table]
    findings: dict[str, int] = field(default_factory=dict)
    # Values outside the tables (comments, metadata...): they do not go into the profile, but they
    # do go into the fingerprint index so the final verification looks for them in the output.
    loose: list[str] = field(default_factory=list)
    # JSON and YAML: the original document, to rebuild it with the same shape
    document: object = None


class UnsupportedFormat(ValueError):
    pass


def read(path: Path) -> Reading:
    suffix = path.suffix.lower().lstrip(".")
    if suffix in {"csv", "tsv"}:
        return _csv(path, suffix)
    if suffix in {"xlsx", "xlsm"}:
        return _xlsx(path, suffix)
    if suffix in STRUCTURED:
        from . import structured

        return structured.read(path, STRUCTURED[suffix])
    raise UnsupportedFormat(suffix)


def _csv(path: Path, file_format: str) -> Reading:
    data = path.read_bytes()
    for encoding in ENCODINGS:
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if file_format == "tsv" else csv.excel
    rows = list(csv.reader(text.splitlines(), dialect))
    if not rows:
        return Reading(file_format, [])
    return Reading(file_format, [Table(None, rows[0], rows[1:])])


def _xlsx(path: Path, file_format: str) -> Reading:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False, keep_vba=False)
    # A second load with the computed values: the output never carries formulas, which can hold
    # literals (="Juan") or references to hidden sheets.
    computed = load_workbook(path, data_only=True, keep_vba=False)
    tables: list[Table] = []
    loose: list[str] = []
    findings = {k: 0 for k in ("hidden_sheets", "comments", "defined_names", "metadata", "headers_footers",
                               "validations", "formulas", "pivot_tables", "charts", "embedded_objects",
                               "external_links")}

    for sheet, sheet_values in zip(workbook.worksheets, computed.worksheets):
        if sheet.sheet_state != "visible":
            findings["hidden_sheets"] += 1
        rows = [list(r) for r in sheet_values.iter_rows(values_only=True)]
        rows = [r for r in rows if any(v is not None for v in r)]
        if rows:
            tables.append(Table(sheet.title, ["" if v is None else str(v) for v in rows[0]], rows[1:],
                                hidden=sheet.sheet_state != "visible"))
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment:
                    findings["comments"] += 1
                    loose += [cell.comment.text, cell.comment.author or ""]
                if cell.data_type == "f":
                    findings["formulas"] += 1
                    loose.append(str(cell.value))
        for part in (sheet.oddHeader, sheet.oddFooter, sheet.evenHeader, sheet.evenFooter):
            for zone in (part.left, part.center, part.right):
                if zone.text:
                    findings["headers_footers"] += 1
                    loose.append(zone.text)
        for validation in sheet.data_validations.dataValidation:
            findings["validations"] += 1
            loose += [validation.formula1 or "", validation.formula2 or ""]
        findings["pivot_tables"] += len(getattr(sheet, "_pivots", []))
        findings["charts"] += len(getattr(sheet, "_charts", []))
        findings["embedded_objects"] += len(getattr(sheet, "_images", []))

    for name, defined in workbook.defined_names.items():
        findings["defined_names"] += 1
        loose += [name, defined.attr_text or ""]
    properties = workbook.properties
    for attribute in ("creator", "lastModifiedBy", "title", "subject", "description", "keywords"):
        if value := getattr(properties, attribute, None):
            findings["metadata"] += 1
            loose.append(str(value))
    findings["external_links"] = len(workbook._external_links)
    return Reading(file_format, tables, findings, [s for s in loose if s])


def write(tables: list[Table], path: Path) -> None:
    """Write a brand-new file: only headers and values, nothing inherited."""
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        if len(tables) != 1:
            raise UnsupportedFormat("csv with several tables")
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle, delimiter="\t" if suffix == ".tsv" else ",")
            writer.writerow([_csv_cell(v) for v in tables[0].columns])
            writer.writerows([_csv_cell(v) for v in row] for row in tables[0].rows)
        return
    if suffix == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.creator = "anonymizer"
        for table in tables:
            sheet = workbook.create_sheet(table.name[:31])
            for row in [table.columns, *table.rows]:
                sheet.append(list(row))
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.data_type = "s"  # text, never a formula
        workbook.save(path)
        return
    raise UnsupportedFormat(suffix.lstrip("."))


def _csv_cell(value: object) -> object:
    # Prevents formula injection when someone opens the CSV in Excel
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        if value[:1] in "+-" and value[1:].replace(".", "", 1).replace(",", "", 1).isdigit():
            return value  # a signed number, not a formula
        return "'" + value
    return "" if value is None else value
