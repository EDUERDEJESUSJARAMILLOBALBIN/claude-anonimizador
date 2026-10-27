import csv
import json

import pytest

from anonymizer import engine, state
from canaries.audit import audit


@pytest.fixture(autouse=True)
def profiled(traps):
    """apply requires a previous profile, as in the real flow."""
    for file in ("clientes.csv", "clientes_latin1.csv", "clientes.xlsx"):
        engine.profile(traps / file)


@pytest.mark.parametrize("file", ["clientes.csv", "clientes_latin1.csv", "clientes.xlsx"])
def test_the_suggested_plan_gives_a_copy_without_canaries(traps, tmp_path, file):
    suffix = ".xlsx" if file.endswith("xlsx") else ".csv"
    result = engine.apply(traps / file, {}, output=tmp_path / "anonimizados" / "out" / f"x.anon{suffix}")
    assert result["verification"]["matches"] == 0
    assert audit(traps / "manifest.json", [tmp_path / "anonimizados" / "out"]) == {}


def test_the_copy_keeps_what_is_useful_and_pseudonymizes_the_id(traps, tmp_path):
    destination = tmp_path / "anonimizados" / "out.csv"
    engine.apply(traps / "clientes.csv", {}, output=destination)
    rows = list(csv.reader(destination.open()))
    assert rows[0] == ["id_cliente", "fecha_nacimiento", "ciudad", "monto"]
    assert rows[1][0].startswith("ID_") and rows[1][1].count("-") == 1
    assert len(rows) == 31


def test_pseudonyms_are_stable_between_runs(traps, tmp_path):
    a = engine.apply(traps / "clientes.csv", {}, output=tmp_path / "anonimizados" / "a.csv")
    b = engine.apply(traps / "clientes.csv", {}, output=tmp_path / "anonimizados" / "b.csv")
    assert (tmp_path / "anonimizados" / "a.csv").read_text() == (tmp_path / "anonimizados" / "b.csv").read_text()
    assert a["tables"] == b["tables"]


def test_a_direct_identifier_cannot_be_kept(traps, tmp_path):
    with pytest.raises(engine.InvalidPlan, match="email"):
        engine.apply(traps / "clientes.csv", {"t1": {"c4": "keep"}}, output=tmp_path / "anonimizados" / "x.csv")
    assert not (tmp_path / "anonimizados" / "x.csv").exists()


def test_kept_free_text_with_phone_numbers_is_rejected_and_deleted(traps, tmp_path):
    with pytest.raises(engine.InvalidPlan, match="verification"):
        engine.apply(traps / "clientes.csv", {"t1": {"c10": "keep"}}, output=tmp_path / "anonimizados" / "x.csv")
    assert not (tmp_path / "anonimizados" / "x.csv").exists()


def test_the_output_cannot_land_on_a_protected_file(traps):
    state.protect(traps / "clientes.csv")
    with pytest.raises(engine.InvalidPlan, match="original"):
        engine.apply(traps / "clientes.csv", {}, output=traps / "clientes.csv")


def test_profiling_protects_the_original_and_saves_fingerprints(traps):
    engine.profile(traps / "clientes.xlsx")
    files, _ = state.protected()
    assert (traps / "clientes.xlsx").resolve() in files
    assert len(state.global_index()) > 100


def test_xlsx_output_has_no_hidden_sheets_or_metadata(traps, tmp_path):
    from openpyxl import load_workbook

    destination = tmp_path / "anonimizados" / "x.xlsx"
    engine.apply(traps / "clientes.xlsx", {"t2": "omit", "t3": "omit"}, output=destination)
    workbook = load_workbook(destination)
    assert [s.sheet_state for s in workbook.worksheets] == ["visible"]
    assert workbook.properties.creator == "anonymizer" and not workbook.defined_names
    assert json.dumps([[c.comment for c in row] for row in workbook.active.iter_rows()]).count("null") > 0


def test_the_settings_block_closes_extensions_and_opens_the_output():
    from anonymizer.settings import block

    filesystem = block()["sandbox"]["filesystem"]
    assert "~/**/*.csv" in filesystem["denyRead"] and "/Volumes/**/*.xlsx" in filesystem["denyRead"]
    # Tested on Seatbelt: the exception only works written with a per-extension wildcard
    assert {"~/anonimizados/*.csv", "~/anonimizados/**/*.csv"} <= set(filesystem["allowRead"])
    assert "~/**/*.db" in filesystem["denyRead"]  # SQLite databases too


# One-command install: merges without deleting, backs up, is idempotent and can be undone
def test_install_merge_and_uninstall(tmp_path, monkeypatch):
    from anonymizer import settings

    home = tmp_path / "home"
    (home / ".claude").mkdir(parents=True)
    settings_file = home / ".claude" / "settings.json"
    own = {"model": "opus", "permissions": {"deny": ["Bash(rm -rf *)"], "defaultMode": "auto"},
           "hooks": {"Stop": []}, "sandbox": {"network": {"allowedDomains": ["github.com"]}}}
    settings_file.write_text(json.dumps(own))
    monkeypatch.setenv("HOME", str(home))
    monkeypatch.delenv("CLAUDE_CONFIG_DIR", raising=False)

    assert "rules added" in settings.install()
    installed = json.loads(settings_file.read_text())
    assert installed["model"] == "opus" and installed["hooks"] == {"Stop": []}
    assert "Bash(rm -rf *)" in installed["permissions"]["deny"] and "Read(~/sensibles/**)" in installed["permissions"]["deny"]
    assert installed["sandbox"]["network"]["allowedDomains"] == ["github.com"] and installed["sandbox"]["enabled"] is True
    assert settings.sandbox_problems(home / "project") == []
    assert list((home / ".claude").glob("settings.json.before-anonymizer-*"))
    assert (home / "sensibles").is_dir() and (home / "anonimizados").is_dir()
    assert "nothing changed" in settings.install()

    assert "Removed" in settings.uninstall()
    assert json.loads(settings_file.read_text()) == own


def test_allow_a_project_folder(tmp_path, monkeypatch):
    from anonymizer import settings

    home = tmp_path / "home"
    project = home / "projects" / "baloto"
    project.mkdir(parents=True)
    monkeypatch.setenv("HOME", str(home))
    monkeypatch.delenv("CLAUDE_CONFIG_DIR", raising=False)
    settings.allow(str(project))
    allow_read = json.loads((home / ".claude" / "settings.json").read_text())["sandbox"]["filesystem"]["allowRead"]
    assert "~/projects/baloto/**/*.csv" in allow_read
    for too_broad in (home, home / "Downloads"):
        too_broad.mkdir(exist_ok=True)
        with pytest.raises(SystemExit):
            settings.allow(str(too_broad))
