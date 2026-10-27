"""Generate trap files with synthetic personal data hidden in every corner.

Every sensitive value is unique and recorded in `manifest.json`. The auditor uses it to know whether
any canary ended up where it should not. This script never prints values: whoever runs it (a person
or Claude) must not see them, so the audit of their own session stays clean.

The data itself is Colombian and in Spanish (names, headers, cities), because that is what the
detector has to recognize in real exports.

Usage: python canaries/generate.py <output_dir> [--rows N] [--seed S]
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import sqlite3
from pathlib import Path

FIRST_NAMES = ["Ximena", "Bernardo", "Leidy", "Yesid", "Marleny", "Duvan", "Rubiela", "Faiver", "Nidia", "Jhon Jairo"]
LAST_NAMES = ["Quintero", "Aristizábal", "Ocampo", "Zuluaga", "Mosquera", "Cárdenas", "Henao", "Rentería", "Galvis", "Úsuga"]
CITIES = ["Medellín", "Pereira", "Quibdó", "Manizales", "Cali"]
HEADER = ["id_cliente", "nombre", "cedula", "correo", "celular", "direccion", "fecha_nacimiento", "ciudad", "monto",
          "observaciones"]
SENSITIVE_COLUMNS = ("nombre", "cedula", "correo", "celular", "direccion", "observaciones")


def records(n: int, rng: random.Random) -> list[dict[str, str]]:
    result = []
    used: set[str] = set()
    for i in range(1, n + 1):
        while True:
            name = f"{rng.choice(FIRST_NAMES)} {rng.choice(LAST_NAMES)} {rng.choice(LAST_NAMES)}"
            if name not in used:
                used.add(name)
                break
        tag = f"{rng.getrandbits(32):08x}"
        result.append(
            {
                "id_cliente": f"CL-{i:05d}",
                "nombre": name,
                "cedula": f"{rng.randint(1_000_000_000, 1_099_999_999):,}".replace(",", "."),
                "correo": f"canary.{tag}@ejemplo.invalid",
                "celular": f"3{rng.randint(0, 5)}{rng.randint(0, 99_999_999):08d}",
                "direccion": f"Calle {rng.randint(1, 120)} # {rng.randint(1, 99)}-{rng.randint(1, 99)} apto {tag[:3]}",
                "fecha_nacimiento": f"{rng.randint(1950, 2005)}-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}",
                "ciudad": rng.choice(CITIES),
                "monto": f"{rng.randint(10_000, 9_000_000)}",
                "observaciones": f"Cliente {name.split()[0]} pidió llamar al {rng.randint(300, 350)}"
                                 f"{rng.randint(1_000_000, 9_999_999)} antes de las 5",
            }
        )
    return result


def sensitive_values(rows: list[dict[str, str]]) -> list[str]:
    return [r[c] for r in rows for c in SENSITIVE_COLUMNS]


def write_csv(rows, path: Path, delimiter: str, encoding: str) -> None:
    with path.open("w", newline="", encoding=encoding) as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADER, delimiter=delimiter)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows, extra: list[str], path: Path) -> dict[str, int]:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "clientes"
    sheet.append(HEADER)
    for r in rows:
        sheet.append([r[c] for c in HEADER])

    hidden = workbook.create_sheet("respaldo")
    hidden.sheet_state = "hidden"
    very_hidden = workbook.create_sheet("sistema")
    very_hidden.sheet_state = "veryHidden"
    for r in rows[:3]:
        hidden.append([r["nombre"], r["cedula"]])
        very_hidden.append([r["correo"]])

    sheet["B2"].comment = Comment(f"Verificar con {extra[0]}", extra[1])
    workbook.defined_names["contacto"] = DefinedName("contacto", attr_text=f'"{extra[2]}"')
    workbook.properties.creator = extra[3]
    workbook.properties.title = extra[4]
    sheet.oddHeader.center.text = extra[5]
    sheet.oddFooter.right.text = extra[6]
    validation = DataValidation(type="list", formula1=f'"{extra[7]},otro"')
    sheet.add_data_validation(validation)
    validation.add("J2")
    sheet["L1"] = f'="{extra[8]}"'
    workbook.save(path)
    return {"hidden_sheets": 2, "comments": 1, "defined_names": 1, "metadata": 2, "headers_footers": 2,
            "validations": 1, "formulas": 1}


def write_sql(rows, extra: list[str], path: Path) -> dict[str, int]:
    def literal(value: str) -> str:
        return "'" + value.replace("'", "''") + "'"

    lines = [
        f"-- exportado por {extra[0]}",
        "CREATE TABLE clientes (" + ", ".join(f"{c} TEXT" for c in HEADER) + ", perfil JSON);",
        f"CREATE VIEW vip AS SELECT * FROM clientes WHERE correo = {literal(extra[1])};",
        "CREATE FUNCTION saludo() RETURNS text AS $$ SELECT " + literal(f"Hola {extra[2]}") + " $$ LANGUAGE sql;",
    ]
    for r in rows[: len(rows) // 2]:
        profile = json.dumps({"contacto": {"alterno": r["correo"], "notas": [r["nombre"]]}}, ensure_ascii=False)
        values = ", ".join(literal(r[c]) for c in HEADER)
        lines.append(f"INSERT INTO clientes VALUES ({values}, {literal(profile)});")
    lines.append(f"COPY clientes ({', '.join(HEADER)}) FROM stdin;")
    for r in rows[len(rows) // 2:]:
        lines.append("\t".join(r[c] for c in HEADER))
    lines.append("\\.")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {"sql_views": 1, "sql_routines": 1, "sql_comments": 1}


def nested_document(rows, extra: list[str]) -> dict:
    """The usual shape of a CRM or API export: records with objects and lists inside."""
    return {
        "exportado_por": extra[9],
        "clientes": [
            {
                "id_cliente": r["id_cliente"],
                "nombre": r["nombre"],
                "contacto": {"correo": r["correo"], "celular": r["celular"], "direccion": r["direccion"]},
                "documentos": [r["cedula"]],
                "ciudad": r["ciudad"],
                "monto": int(r["monto"]),
                "notas": [{"texto": r["observaciones"]}],
            }
            for r in rows
        ],
    }


def write_structured(rows, extra: list[str], folder: Path) -> None:
    import yaml

    document = nested_document(rows, extra)
    (folder / "clientes.json").write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")
    (folder / "clientes.yaml").write_text(yaml.safe_dump(document, allow_unicode=True, sort_keys=False),
                                          encoding="utf-8")


def write_sqlite(rows, path: Path) -> None:
    path.unlink(missing_ok=True)
    with sqlite3.connect(path) as connection:
        connection.execute("CREATE TABLE clientes (" + ", ".join(f"{c} TEXT" for c in HEADER) + ")")
        connection.executemany(
            f"INSERT INTO clientes VALUES ({', '.join('?' for _ in HEADER)})",
            [[r[c] for c in HEADER] for r in rows],
        )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("output", type=Path)
    parser.add_argument("--rows", type=int, default=50)
    parser.add_argument("--seed", type=int, default=None)
    args = parser.parse_args()

    rng = random.Random(args.seed)
    args.output.mkdir(parents=True, exist_ok=True)
    rows = records(args.rows, rng)
    extra = [f"canary.extra{i}.{rng.getrandbits(32):08x}@ejemplo.invalid" for i in range(12)]

    write_csv(rows, args.output / "clientes.csv", ",", "utf-8")
    write_csv(rows, args.output / "clientes_latin1.csv", ";", "latin-1")
    hidden = {
        "clientes.xlsx": write_xlsx(rows, extra, args.output / "clientes.xlsx"),
        "clientes.sql": write_sql(rows, extra, args.output / "clientes.sql"),
    }
    write_sqlite(rows, args.output / "clientes.sqlite")
    write_structured(rows, extra, args.output)

    manifest = {"values": sorted(set(sensitive_values(rows) + extra)), "hidden": hidden}
    (args.output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False))
    files = sorted(p.name for p in args.output.iterdir() if p.name != "manifest.json")
    print(f"{len(manifest['values'])} canaries in {len(files)} files: {', '.join(files)}")


if __name__ == "__main__":
    main()
